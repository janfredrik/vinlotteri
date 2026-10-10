#!/usr/bin/env python3
"""
One-off: read the "Vinlotteri loddark - <year>.xlsx" (one sheet per week) and
emit scripts/loddark-<year>.json for import-loddark.mjs to load into the DB.

Needs openpyxl:  python3 -m venv .venv && .venv/bin/pip install openpyxl
Run:             .venv/bin/python scripts/extract-loddark.py "Vinlotteri loddark - 2026.xlsx"
"""
import json
import os
import re
import sys
from datetime import date

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Vinlotteri loddark - 2026.xlsx"
YEAR = int(re.search(r"(20\d\d)", SRC).group(1))
OUT = f"scripts/loddark-{YEAR}.json"

MONEY_LABELS = {
    "Inn fra forrige gang": "carried_over",
    "Lodd solgt": "lodd_solgt",
    "Kost på polet": "polet_cost",
    "Totalt videreført": "videre",
}


def find_grid(ws):
    # The 1..6 header sits in B1 in most sheets, but some weeks moved it (D7, D14, ...).
    for row in ws.iter_rows():
        for cell in row:
            r, c = cell.row, cell.column
            if [ws.cell(row=r, column=c + i).value for i in range(6)] != [1, 2, 3, 4, 5, 6]:
                continue
            cells = []
            for gr in range(1, 7):
                for gc in range(1, 7):
                    v = ws.cell(row=r + gr, column=c - 1 + gc).value
                    if v is not None and str(v).strip():
                        cells.append({"row": gr, "col": gc, "name": str(v).strip()})
            return cells
    return []


def find_money(ws):
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            key = MONEY_LABELS.get(str(cell.value).strip()) if isinstance(cell.value, str) else None
            if key:
                out[key] = ws.cell(row=cell.row, column=cell.column + 3).value
    return out


def find_text(ws, prefix):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower().startswith(prefix):
                return cell.value.strip()
    return None


def parse_vipps(text):
    if not text:
        return "", ""
    body = re.sub(r"^VIPPS:?\s*", "", text, flags=re.I)
    name = ""
    m = re.search(r"\(([^)]+)\)", body)
    if m:
        name = m.group(1).strip()
        body = body.replace(m.group(0), " ")
    digits = re.findall(r"[\d ]{6,}", body)
    number = re.sub(r"\D", "", digits[0]) if digits else ""
    if not name:
        m = re.search(r"[-–]\s*(.+)$", body)
        if m:
            name = m.group(1).strip()
    return name, number


def parse_winners(text):
    if not text:
        return []
    text = re.sub(r"^vinnere:?\s*", "", text.strip(), flags=re.I)
    marks = list(re.finditer(r"(\d+)\s*[:.]\s*", text))
    out = []
    for i, m in enumerate(marks):
        rank = int(m.group(1))
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        body = text[m.end():end].strip().strip("-–").strip()
        if not body:
            continue
        parts = re.split(r"\s*[-–]\s*|\s*:\s+", body, maxsplit=1)
        name = parts[0].strip()
        prize = parts[1].strip() if len(parts) > 1 else None
        out.append({"rank": rank, "name": name, "prize": prize})
    return out


wb = openpyxl.load_workbook(SRC, data_only=True)
rounds = []
for name in wb.sheetnames:
    wk = re.search(r"(\d+)", name)
    if not wk:
        continue
    week = int(wk.group(1))
    ws = wb[name]
    money = find_money(ws)
    org, number = parse_vipps(find_text(ws, "vipps"))
    winners = parse_winners(find_text(ws, "vinnere"))
    try:
        d = date.fromisocalendar(YEAR, week, 4)  # Thursday
    except ValueError:
        d = date.fromisocalendar(YEAR, week, 1)
    rounds.append(
        {
            "week": week,
            "title": f"Uke {week} – {YEAR}",
            "date": d.isoformat(),
            "ticket_price": 15,
            "carried_over": money.get("carried_over"),
            "polet_cost": money.get("polet_cost"),
            "sheet_lodd_solgt": money.get("lodd_solgt"),
            "sheet_videre": money.get("videre"),
            "organizer_name": org,
            "vipps_number": number,
            "grid_size": 6,
            "cells": find_grid(ws),
            "winners": winners,
        }
    )

rounds.sort(key=lambda r: r["week"])
with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"source": os.path.basename(SRC), "year": YEAR, "rounds": rounds}, f, ensure_ascii=False, indent=2)

for r in rounds:
    sold = len(r["cells"])
    calc = (r["carried_over"] or 0) + sold * 15 - (r["polet_cost"] or 0)
    flag = "" if abs(calc - (r["sheet_videre"] or 0)) < 0.5 else "  <-- videre mismatch"
    print(f"Uke {r['week']:2}  {sold:2} lodd  org={r['organizer_name']!r:20} {r['vipps_number']:9} "
          f"winners={len(r['winners'])}  videre calc={calc} sheet={r['sheet_videre']}{flag}")
print(f"\nwrote {OUT}")
